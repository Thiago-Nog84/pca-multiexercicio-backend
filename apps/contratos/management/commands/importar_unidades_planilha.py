"""
Management command: importar_unidades_planilha
===============================================
Lê a planilha 'contratos_sem_unidade.xlsx' preenchida manualmente
(coluna 'Unidade') e atualiza Contrato.unidade_requisitante no banco.

Uso:
  python manage.py importar_unidades_planilha --dry-run
  python manage.py importar_unidades_planilha
"""

import os
import re

from django.core.management.base import BaseCommand, CommandError

# Correções de typos/variantes conhecidas
_CORRECOES = {
    "CCA": "CCS",   # typo em 25017488 (confecção de vestimentas → CCS)
    "C.T.I": "CTI",
    "FPROCON": "FPROCON",
}


class Command(BaseCommand):
    help = "Importa vínculos contrato→unidade da planilha preenchida manualmente"

    def add_arguments(self, parser):
        default = os.path.normpath(
            os.path.join(
                os.path.dirname(__file__),
                "..", "..", "..", "..", "data", "contratos",
                "contratos_sem_unidade.xlsx",
            )
        )
        parser.add_argument("--planilha", default=default,
                            help="Caminho da planilha (padrão: data/contratos/contratos_sem_unidade.xlsx)")
        parser.add_argument("--dry-run", action="store_true",
                            help="Mostra o que seria feito sem salvar")
        parser.add_argument("--force", action="store_true",
                            help="Sobrescreve vínculos já existentes")

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("Instale openpyxl: pip install openpyxl")

        from apps.contratos.models import Contrato
        from apps.core.models import UnidadeRequisitante

        dry_run = options["dry_run"]
        force   = options["force"]
        path    = options["planilha"]

        if not os.path.exists(path):
            raise CommandError(f"Arquivo não encontrado: {path}")

        if dry_run:
            self.stdout.write(self.style.WARNING("⚠ DRY-RUN — nenhuma alteração será salva\n"))

        # ── Ler planilha ──────────────────────────────────────────────────────
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

        if "Nº Contrato" not in headers or "Unidade" not in headers:
            raise CommandError(
                "Planilha não tem as colunas esperadas ('Nº Contrato' e 'Unidade'). "
                f"Colunas encontradas: {headers}"
            )

        # Montar lista de (numero_contrato, sigla_unidade)
        entradas: list[tuple[str, str]] = []
        for r in range(2, ws.max_row + 1):
            row = {headers[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
            num = str(row.get("Nº Contrato") or "").strip()
            uni = str(row.get("Unidade") or "").strip().upper()
            uni = _CORRECOES.get(uni, uni)
            if num and uni and uni not in ("NONE", "—", ""):
                entradas.append((num, uni))

        self.stdout.write(f"Planilha: {len(entradas)} linhas com unidade preenchida\n")

        # ── Cache de unidades ──────────────────────────────────────────────────
        unidades = {u.sigla.upper(): u for u in UnidadeRequisitante.objects.all()}

        # ── Processar ─────────────────────────────────────────────────────────
        atualizados = 0
        nao_encontrados_contrato: list[str] = []
        nao_encontrados_unidade:  list[tuple[str, str]] = []
        ja_vinculados = 0

        for num_contrato, sigla in entradas:
            # Buscar contrato: match exato primeiro, depois por número normalizado
            contrato = Contrato.objects.filter(numero_contrato=num_contrato).first()
            if not contrato:
                # Tentar match normalizado (sem sufixo de órgão)
                m = re.search(r"(\d+)/((19|20)\d{2})", num_contrato)
                if m:
                    seq  = m.group(1).lstrip("0") or "0"
                    ano  = m.group(2)
                    num_norm = f"{seq}/{ano}"
                    # Buscar contratos cujo numero_contrato normalizado bate
                    for c in Contrato.objects.filter(numero_contrato__contains=f"/{ano}"):
                        m2 = re.search(r"(\d+)/((19|20)\d{2})", c.numero_contrato)
                        if m2:
                            seq2 = m2.group(1).lstrip("0") or "0"
                            if seq2 == seq and m2.group(2) == ano:
                                contrato = c
                                break

            if not contrato:
                nao_encontrados_contrato.append(num_contrato)
                continue

            if contrato.unidade_requisitante and not force:
                ja_vinculados += 1
                continue

            unidade = unidades.get(sigla)
            if not unidade:
                nao_encontrados_unidade.append((num_contrato, sigla))
                continue

            atualizados += 1
            if dry_run:
                self.stdout.write(
                    f"  [DRY] {num_contrato:<30} → {sigla}"
                    + (f"  (era: {contrato.unidade_requisitante.sigla})" if contrato.unidade_requisitante else "")
                )
            else:
                contrato.unidade_requisitante = unidade
                contrato.save(update_fields=["unidade_requisitante"])
                self.stdout.write(f"  ✓ {num_contrato:<30} → {sigla}")

        # ── Resumo ─────────────────────────────────────────────────────────────
        self.stdout.write("\n" + "─" * 60)
        self.stdout.write(self.style.SUCCESS(
            f"{'[DRY-RUN] ' if dry_run else ''}Atualizados: {atualizados} contratos"
        ))
        if ja_vinculados:
            self.stdout.write(
                f"  Ignorados (já tinham vínculo — use --force): {ja_vinculados}"
            )
        if nao_encontrados_contrato:
            self.stdout.write(self.style.WARNING(
                f"\nContratos não encontrados no banco ({len(nao_encontrados_contrato)}):"
            ))
            for n in nao_encontrados_contrato:
                self.stdout.write(f"  {n}")
        if nao_encontrados_unidade:
            self.stdout.write(self.style.WARNING(
                f"\nSiglas sem UnidadeRequisitante no banco ({len(nao_encontrados_unidade)}):"
            ))
            for n, s in nao_encontrados_unidade:
                self.stdout.write(f"  {n} → '{s}'")

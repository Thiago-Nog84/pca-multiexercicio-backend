"""
Management command: exportar_sem_unidade
=========================================
Exporta todos os contratos sem unidade_requisitante para um arquivo Excel
com categorização automática e sugestão de fonte de dados.

Uso:
  python manage.py exportar_sem_unidade
  python manage.py exportar_sem_unidade --output caminho/arquivo.xlsx
"""

import os
import re

from django.core.management.base import BaseCommand


def _categorizar(numero: str) -> tuple[str, str]:
    """
    Retorna (categoria, fonte_sugerida) com base no padrão do numero_contrato.
    """
    n = (numero or "").strip()

    if re.match(r"^20\d{2}N[ER]\d+", n):
        return (
            "NE (Nota de Empenho)",
            "Planilha de empenhos 2025 com coluna de setor, ou SIAFE API por número de NE",
        )

    if re.match(r"^\d{8}$", n):
        ano = int(n[:2]) + 2000
        return (
            f"Código SIAFE interno ({ano})",
            f"Planilha de contratos {ano} completa, ou SIAFE API por codContrato",
        )

    m = re.search(r"\d+/((19|20)\d{2})", n)
    if m:
        ano = int(m.group(1))
        if ano <= 2019:
            return (
                f"Contrato antigo ({ano})",
                "Sem planilha disponível — vinculação manual no admin",
            )
        if ano in (2021, 2022):
            return (
                f"Contrato {ano} (planilha incompleta)",
                f"Planilha PGJ/FMMPPI de {ano} (arquivo atual cobre apenas FPDC)",
            )
        return (
            f"Contrato {ano} sem setor na planilha",
            f"Verificar planilha de processos de {ano}",
        )

    if re.match(r"^\d{5,7}$", n):
        return (
            "Código numérico SIAFE (5-7 dígitos)",
            "Planilha do exercício correspondente ou SIAFE API",
        )

    return (
        "Formato desconhecido",
        "Vinculação manual no admin",
    )


class Command(BaseCommand):
    help = "Exporta contratos sem unidade requisitante para Excel"

    def add_arguments(self, parser):
        default_out = os.path.normpath(
            os.path.join(
                os.path.dirname(__file__),
                "..", "..", "..", "..", "data", "contratos",
                "contratos_sem_unidade.xlsx",
            )
        )
        parser.add_argument(
            "--output",
            default=default_out,
            help=f"Caminho do arquivo Excel de saída (padrão: {default_out})",
        )

    def handle(self, *args, **options):
        try:
            import openpyxl
            from openpyxl.styles import (
                Alignment, Border, Font, PatternFill, Side
            )
            from openpyxl.utils import get_column_letter
        except ImportError:
            self.stderr.write("Instale openpyxl: pip install openpyxl")
            return

        from apps.contratos.models import Contrato

        contratos = (
            Contrato.objects
            .filter(unidade_requisitante__isnull=True)
            .select_related("orgao", "gestor")
            .order_by("numero_contrato")
        )
        total = contratos.count()
        self.stdout.write(f"Exportando {total} contratos sem unidade requisitante...")

        # ── Workbook ──────────────────────────────────────────────────────────
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sem Unidade"

        # Estilos
        font_title = Font(name="Arial", bold=True, size=13, color="FFFFFF")
        font_header = Font(name="Arial", bold=True, size=10, color="FFFFFF")
        font_data   = Font(name="Arial", size=9)
        font_cat    = Font(name="Arial", size=9, italic=True, color="4A4A4A")

        fill_title  = PatternFill("solid", start_color="1F3864")
        fill_header = PatternFill("solid", start_color="2E5A9C")
        fill_ne     = PatternFill("solid", start_color="FFF2CC")  # NE
        fill_siafe  = PatternFill("solid", start_color="E2EFDA")  # SIAFE
        fill_antigo = PatternFill("solid", start_color="FCE4D6")  # Antigo/manual
        fill_incomp = PatternFill("solid", start_color="DDEBF7")  # incompleto
        fill_zebra  = PatternFill("solid", start_color="F5F5F5")

        thin = Side(style="thin", color="CCCCCC")
        borda = Border(left=thin, right=thin, top=thin, bottom=thin)

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
        align_right  = Alignment(horizontal="right",  vertical="center")

        # ── Título ────────────────────────────────────────────────────────────
        ws.merge_cells("A1:L1")
        ws["A1"] = f"Contratos sem Unidade Requisitante — {total} registros"
        ws["A1"].font = font_title
        ws["A1"].fill = fill_title
        ws["A1"].alignment = align_center
        ws.row_dimensions[1].height = 28

        # ── Cabeçalhos ────────────────────────────────────────────────────────
        headers = [
            ("Nº Contrato",      18),
            ("Órgão",             9),
            ("Status",           10),
            ("Tipo",             14),
            ("Contratado",       32),
            ("Objeto",           40),
            ("Data Assinatura",  14),
            ("Vigência Fim",     12),
            ("Valor Atual (R$)", 16),
            ("Nº SEI",           22),
            ("Categoria",        22),
            ("Fonte Sugerida",   38),
        ]

        for col, (label, width) in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col, value=label)
            cell.font    = font_header
            cell.fill    = fill_header
            cell.border  = borda
            cell.alignment = align_center
            ws.column_dimensions[get_column_letter(col)].width = width

        ws.row_dimensions[2].height = 32
        ws.freeze_panes = "A3"

        # ── Dados ─────────────────────────────────────────────────────────────
        row_idx = 3
        for i, c in enumerate(contratos):
            categoria, fonte = _categorizar(c.numero_contrato)

            # Cor de fundo por categoria
            if "NE" in categoria:
                bg = fill_ne
            elif "SIAFE" in categoria or "numérico" in categoria:
                bg = fill_siafe
            elif "antigo" in categoria.lower() or "desconhecido" in categoria.lower():
                bg = fill_antigo
            elif "incompleta" in categoria.lower():
                bg = fill_incomp
            else:
                bg = fill_zebra if i % 2 == 0 else PatternFill()

            valores = [
                c.numero_contrato,
                c.orgao.sigla if c.orgao else "",
                c.get_status_display(),
                c.get_tipo_display(),
                c.contratado_razao_social,
                c.objeto[:200],
                c.data_assinatura.strftime("%d/%m/%Y") if c.data_assinatura else "",
                c.data_fim_vigencia.strftime("%d/%m/%Y") if c.data_fim_vigencia else "",
                float(c.valor_atual) if c.valor_atual else 0,
                c.numero_sei or "",
                categoria,
                fonte,
            ]

            for col, val in enumerate(valores, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border = borda
                if col == 11:  # categoria
                    cell.font = font_cat
                elif col == 9:  # valor
                    cell.number_format = '#,##0.00'
                    cell.alignment = align_right
                    cell.font = font_data
                    cell.fill = bg
                    continue
                else:
                    cell.font = font_data
                cell.fill = bg
                cell.alignment = align_left if col in (5, 6, 12) else align_center

            ws.row_dimensions[row_idx].height = 18
            row_idx += 1

        # ── Aba de legenda ────────────────────────────────────────────────────
        leg = wb.create_sheet("Legenda")
        leg.column_dimensions["A"].width = 6
        leg.column_dimensions["B"].width = 28
        leg.column_dimensions["C"].width = 50

        leg.merge_cells("A1:C1")
        leg["A1"] = "Legenda de Cores"
        leg["A1"].font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
        leg["A1"].fill = fill_title
        leg["A1"].alignment = align_center

        legendas = [
            (fill_ne,     "NE (Nota de Empenho)",
             "Cadastrado com número de NE em vez de número de contrato"),
            (fill_siafe,  "Código SIAFE interno",
             "Número automático do SIAFE-PI (8 dígitos)"),
            (fill_antigo, "Contrato antigo / manual",
             "Sem planilha disponível — vinculação manual no admin"),
            (fill_incomp, "Planilha incompleta",
             "Planilha carregada cobre apenas parte do exercício (ex: só FPDC)"),
            (fill_zebra,  "Outros",
             "Demais casos sem correspondência"),
        ]

        for r, (fill, nome, desc) in enumerate(legendas, 2):
            leg.cell(r, 1).fill = fill
            leg.cell(r, 1).border = borda
            leg.cell(r, 2, value=nome).font = Font(name="Arial", bold=True, size=10)
            leg.cell(r, 2).fill = fill
            leg.cell(r, 2).border = borda
            leg.cell(r, 3, value=desc).font = Font(name="Arial", size=10)
            leg.cell(r, 3).fill = fill
            leg.cell(r, 3).border = borda
            leg.row_dimensions[r].height = 20

        # ── Salvar ────────────────────────────────────────────────────────────
        output = options["output"]
        os.makedirs(os.path.dirname(output), exist_ok=True)
        wb.save(output)
        self.stdout.write(self.style.SUCCESS(f"\nArquivo gerado: {output}"))

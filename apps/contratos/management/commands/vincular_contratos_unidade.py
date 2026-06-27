"""
Management command: vincular_contratos_unidade
===============================================
Cruza as planilhas de acompanhamento de contratos e de processos de compras
para vincular automaticamente cada Contrato à sua UnidadeRequisitante.

Fontes suportadas (todas com default em data/contratos/):
  --contratos-AAAA  ACOMPANHAMENTO-CONTRATOS AAAA (.xlsx/.csv)
  --processos-2025  2025 - ACOMPANHAMENTO-COMPRAS.xlsx
  --processos-2026  2026 - ACOMPANHAMENTO DOS PROCESSOS-2026.xlsx

Lógica de resolução (por prioridade):
  1. numero_contrato → setor  (direto, planilhas com coluna SETOR)
  2. numero_sei do banco → setor  (via mapa SEI→setor)
  3. numero_contrato → SEI (planilhas 2025/2026) → setor

Uso:
  python manage.py vincular_contratos_unidade --dry-run
  python manage.py vincular_contratos_unidade
  python manage.py vincular_contratos_unidade --force
"""

import csv
import logging
import os
import re

from django.core.management.base import BaseCommand, CommandError

logger = logging.getLogger(__name__)

# ── Normalização ──────────────────────────────────────────────────────────────

def _norm_sei(sei: str) -> str:
    return re.sub(r"\s+", "", sei or "").upper()


def _norm_numero(num: str) -> str:
    """
    Extrai NN/AAAA de qualquer variante:
      '08/2025/FPDC' → '8/2025'
      'Contrato nº 68/2025' → '68/2025'
      '049/2025' → '49/2025'
    """
    num = (num or "").strip()
    m = re.search(r"(\d+)/((19|20)\d{2})", num)
    if m:
        seq = m.group(1).lstrip("0") or "0"
        return f"{seq}/{m.group(2)}"
    return num.upper()


# Mapa de variações históricas de nome de setor → sigla canônica
_SETOR_ALIASES: dict[str, str | None] = {
    # ── Siglas diretas ──────────────────────────────────────────────
    "CAA": "CAA", "CTI": "CTI", "CPPT": "CPPT", "CRH": "CRH",
    "GSI": "GSI", "CEAF": "CEAF", "CI": "CI", "CLC": "CLC",
    "CCF": "CCF", "APG": "APG", "CONINT": "CONINT",
    "GAECO": "GAECO", "FPROCON": "FPROCON", "CCS": "CCS",
    # ── Variantes com ponto ─────────────────────────────────────────
    "C.T.I": "CTI", "C.T.I.": "CTI",
    "C.L.C": "CLC",
    # ── Nomes longos 2023/2024 ──────────────────────────────────────
    "CAA - COORDENADORIA DE APOIO ADMINISTRATIVO": "CAA",
    "CTI - COORDENADORIA DE TECNOLOGIA DA INFORMAÇÃO": "CTI",
    # ── Nomes longos 2020–2022 ──────────────────────────────────────
    "COORD. APOIO ADM. / DIVISÃO DE MATERIAL PERMANENTE": "CAA",
    "COORD. APOIO ADM./DIVISÃO DE MATERIAL PERMANENTE": "CAA",
    "COORD. DE APOIO ADMINISTRATIVO": "CAA",
    "COORD. APOIO ADMINISTRATIVO": "CAA",
    "COORD. APOIO ADMINISTRATIVO/DIVISÃO MATERIAL PERMANENTE": "CAA",
    "COORD. APOIO ADMINISTRATIVO/DIVISÃO MATERIAL PERMnente": "CAA",
    "DIVISÃO DE MATERIAL PERMANENTE": "CAA",
    "DIVISÃO DE TRANSPORTES": "CAA",
    "COORD. TECNOLOGIA DA INFORMAÇÃO": "CTI",
    "COORD. DE TECNOLOGIA DA INFORMAÇÃO": "CTI",
    "COORD. PERÍCIAS E PARECERES TÉCNICOS": "CPPT",
    "COORD. PERICIA E PARECERES TÉCNICOS": "CPPT",
    "COORD. DE PERÍCIAS E PARECERES TÉCNICOS": "CPPT",
    "COORD. DE RECURSOS HUMANOS": "CRH",
    "COORD. RECURSOS HUMANOS": "CRH",
    "COORD. DE INFRAESTRUTURA": "CI",
    "COORD. INFRAESTRUTURA": "CI",
    "COORD. DE COMUNICAÇÃO SOCIAL": "CCS",
    "COORD. COMUNICAÇÃO SOCIAL": "CCS",
    "GERÊNCIA DE SEGURANÇA INSTITUCIONAL": "GSI",
    "GERENCIA DE SEGURANÇA INSTITUCIONAL": "GSI",
    # ── PROCON/FPROCON ──────────────────────────────────────────────
    "PROCON": "FPROCON",
    "PROCON-MP/PI": "FPROCON",
    "FUNDO PROCON": "FPROCON",
    # ── Aliases compostos ────────────────────────────────────────────
    "CCS/CAA": "CAA",
    "INOVA/LAB": "CTI",
    "LABORATÓRIO DE INOVAÇÃO DE TECNOLOGIA INOVA / MPPI": "CTI",
    # ── Ignorar ─────────────────────────────────────────────────────
    "ASSESPPLAGES": None,
    "PGJ": None,
    "PGPI": None,
    "": None,
}


def _norm_setor(setor: str) -> str | None:
    """Resolve nome/sigla de setor para a sigla canônica de UnidadeRequisitante."""
    s = (setor or "").strip()
    # 1. Match exato no alias map (case-sensitive preservado, mas comparação upper)
    s_up = s.upper()
    for key, val in _SETOR_ALIASES.items():
        if key.upper() == s_up:
            return val
    # 2. Match parcial: pegar sigla no início antes do primeiro " - " ou "/"
    partes = re.split(r"\s*[-/]\s*", s_up)
    candidato = partes[0].strip().rstrip(".")
    if candidato in _SETOR_ALIASES:
        return _SETOR_ALIASES[candidato]
    # 3. Sigla reconhecida em qualquer posição
    siglas_conhecidas = {k for k in _SETOR_ALIASES if len(k) <= 8 and k.isalpha()}
    for sigla in siglas_conhecidas:
        if re.search(r"\b" + sigla + r"\b", s_up):
            return _SETOR_ALIASES[sigla]
    return s or None


# ── Função auxiliar: encontrar linha de cabeçalho ─────────────────────────────

def _find_header(ws, required: list[str]) -> list[str]:
    """Retorna lista de headers da primeira linha que contenha todos os termos."""
    for row in ws.iter_rows(values_only=True):
        vals = [str(v or "").strip() for v in row]
        joined = " ".join(vals).upper()
        if all(t.upper() in joined for t in required):
            return vals
    return []


# ── Leitores ──────────────────────────────────────────────────────────────────

def _ler_contratos_com_setor_xlsx(path: str) -> tuple[dict, dict]:
    """
    Lê planilhas que têm SETOR na mesma linha do contrato (2020–2024).
    Retorna:
      num_setor  → {numero_norm: sigla_setor}
      ne_setor   → {NE_norm: sigla_setor}   (quando há coluna de NE)
    """
    try:
        import openpyxl
    except ImportError:
        raise CommandError("Instale openpyxl: pip install openpyxl")

    num_setor: dict[str, str] = {}
    ne_setor: dict[str, str] = {}

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = _find_header(ws, ["CONTRATO", "SETOR"])
    if not headers:
        return num_setor, ne_setor

    col_num   = next((h for h in headers if "CONTRATO" in h.upper()), None)
    col_setor = next((h for h in headers if "SETOR" in h.upper()), None)
    col_ne    = next(
        (h for h in headers
         if "EMPENHO" in h.upper() and "DATA" not in h.upper()),
        None
    )

    for row in ws.iter_rows(values_only=True):
        vals = [str(v or "").strip() for v in row]
        d = dict(zip(headers, vals))

        num_raw   = (d.get(col_num, "") if col_num else "").strip()
        setor_raw = (d.get(col_setor, "") if col_setor else "").strip()
        ne_raw    = (d.get(col_ne, "") if col_ne else "").strip()

        # Ignorar linhas de cabeçalho ou vazias
        if not num_raw or not re.search(r"\d+/\d{4}", num_raw):
            continue

        setor = _norm_setor(setor_raw)
        if setor:
            num_setor[_norm_numero(num_raw)] = setor

        # Mapear NE → setor (pode ter múltiplas NEs separadas por vírgula)
        if setor and ne_raw:
            for ne in re.split(r"[;,/\s]+", ne_raw):
                ne = ne.strip()
                if re.match(r"20\d{2}N[ER]\d+", ne):
                    ne_setor[ne.upper()] = setor

    return num_setor, ne_setor


def _ler_processos_xlsx(path: str) -> dict:
    """Lê planilhas de processos. Retorna {SEI_norm: sigla_setor}."""
    try:
        import openpyxl
    except ImportError:
        raise CommandError("Instale openpyxl: pip install openpyxl")

    resultado: dict[str, str] = {}
    wb = openpyxl.load_workbook(path, data_only=True)

    for sname in wb.sheetnames:
        ws = wb[sname]
        headers = _find_header(ws, ["SETOR"])
        if not headers:
            continue
        col_sei = next(
            (h for h in headers if "PROCESSO" in h.upper() or h.upper() == "PGEA"),
            None
        )
        col_setor = next((h for h in headers if "SETOR" in h.upper()), None)
        if not col_sei or not col_setor:
            continue

        for row in ws.iter_rows(values_only=True):
            vals = [str(v or "").strip() for v in row]
            d = dict(zip(headers, vals))
            sei   = d.get(col_sei, "").strip()
            setor = _norm_setor(d.get(col_setor, "").strip())
            if sei and setor and re.search(r"\d", sei):
                resultado[_norm_sei(sei)] = setor

    return resultado


def _ler_contratos_2025_csv(path: str) -> dict:
    """Retorna {numero_norm: SEI_norm}."""
    resultado: dict[str, str] = {}
    with open(path, encoding="utf-8-sig") as f:
        lines = f.readlines()
    header_idx = next(
        (i for i, l in enumerate(lines) if "CONTRATO" in l.upper() and "PROCESSO" in l.upper()),
        None
    )
    if header_idx is None:
        return resultado
    for row in csv.DictReader(lines[header_idx:], delimiter=";"):
        num = row.get("Nº CONTRATO", "").strip()
        sei = row.get("Nº PROCESSO", "").strip()
        if num and sei and re.search(r"\d+/\d{4}", num):
            resultado[_norm_numero(num)] = _norm_sei(sei)
    return resultado


def _ler_contratos_2026_xlsx(path: str) -> dict:
    """Retorna {numero_norm: SEI_norm}."""
    try:
        import openpyxl
    except ImportError:
        raise CommandError("Instale openpyxl: pip install openpyxl")

    resultado: dict[str, str] = {}
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = _find_header(ws, ["PROCESSO", "CONTRATO"])
    if not headers:
        return resultado

    for row in ws.iter_rows(values_only=True):
        vals = [str(v or "").strip() for v in row]
        d = dict(zip(headers, vals))
        num = d.get("Nº CONTRATO", "") or next((v for v in vals if re.match(r"^\d{1,3}/20\d{2}$", v)), "")
        sei = d.get("Nº PROCESSO", "") or next((v for v in vals if re.match(r"^19\.21\.", v)), "")
        if num and sei and re.search(r"\d+/\d{4}", num):
            resultado[_norm_numero(num)] = _norm_sei(sei)
    return resultado


# ── Command ───────────────────────────────────────────────────────────────────

class Command(BaseCommand):
    help = "Vincula contratos às unidades requisitantes via cruzamento de planilhas"

    def add_arguments(self, parser):
        base = os.path.normpath(
            os.path.join(os.path.dirname(__file__), "..", "..", "..", "..", "data", "contratos")
        )
        for ano in ["2020", "2021", "2022", "2023", "2024"]:
            parser.add_argument(
                f"--contratos-{ano}",
                default=os.path.join(base, f"ACOMPANHAMENTO-CONTRATOS {ano}.xlsx"),
                help=f"Planilha {ano} com coluna SETOR (xlsx)",
            )
        parser.add_argument(
            "--contratos-2025",
            default=os.path.join(base, "ACOMPANHAMENTO-CONTRATOS 2025.csv"),
        )
        parser.add_argument(
            "--contratos-2026",
            default=os.path.join(base, "ACOMPANHAMENTO-CONTRATOS 2026.xlsx"),
        )
        parser.add_argument(
            "--processos-2025",
            default=os.path.join(base, "2025 - ACOMPANHAMENTO-COMPRAS.xlsx"),
        )
        parser.add_argument(
            "--processos-2026",
            default=os.path.join(base, "2026 - ACOMPANHAMENTO DOS PROCESSOS-2026.xlsx"),
        )
        parser.add_argument("--dry-run", action="store_true")
        parser.add_argument("--force", action="store_true",
                            help="Sobrescreve vínculos já existentes")

    def handle(self, *args, **options):
        from apps.contratos.models import Contrato
        from apps.core.models import UnidadeRequisitante

        dry_run = options["dry_run"]
        force   = options["force"]

        if dry_run:
            self.stdout.write(self.style.WARNING("⚠ DRY-RUN — nenhuma alteração será salva\n"))

        # ── 1. Mapa direto numero→setor  e  NE→setor  (planilhas 2020–2024) ──
        num_setor: dict[str, str] = {}
        ne_setor:  dict[str, str] = {}

        self.stdout.write("Carregando planilhas com setor direto (2020–2024)...")
        for ano in ["2020", "2021", "2022", "2023", "2024"]:
            path = options[f"contratos_{ano}"]
            if not os.path.exists(path):
                self.stdout.write(self.style.WARNING(f"  [{ano}] não encontrado"))
                continue
            n_s, ne_s = _ler_contratos_com_setor_xlsx(path)
            num_setor.update(n_s)
            ne_setor.update(ne_s)
            self.stdout.write(
                f"  [{ano}] {len(n_s)} contrato→setor  |  {len(ne_s)} NE→setor"
            )

        # ── 2. Mapa SEI→setor (processos 2025/2026) ──────────────────────────
        sei_setor: dict[str, str] = {}

        self.stdout.write("\nCarregando planilhas SEI→setor...")
        for label, path in [
            ("processos-2025", options["processos_2025"]),
            ("processos-2026", options["processos_2026"]),
        ]:
            if not os.path.exists(path):
                self.stdout.write(self.style.WARNING(f"  [{label}] não encontrado"))
                continue
            data = _ler_processos_xlsx(path)
            sei_setor.update(data)
            self.stdout.write(f"  [{label}] {len(data)} SEI→setor")

        # ── 3. Mapa numero→SEI (planilhas 2025/2026) ─────────────────────────
        num_sei: dict[str, str] = {}

        self.stdout.write("\nCarregando planilhas contrato→SEI (2025/2026)...")
        for label, path, reader in [
            ("contratos-2025", options["contratos_2025"], _ler_contratos_2025_csv),
            ("contratos-2026", options["contratos_2026"], _ler_contratos_2026_xlsx),
        ]:
            if not os.path.exists(path):
                self.stdout.write(self.style.WARNING(f"  [{label}] não encontrado"))
                continue
            data = reader(path)
            num_sei.update(data)
            self.stdout.write(f"  [{label}] {len(data)} contrato→SEI")

        # ── 4. Cache de unidades ──────────────────────────────────────────────
        unidades = {u.sigla.upper(): u for u in UnidadeRequisitante.objects.all()}
        if not unidades:
            raise CommandError("Nenhuma UnidadeRequisitante no banco.")
        self.stdout.write(f"\nUnidades: {sorted(unidades.keys())}\n")

        # ── 5. Processar contratos ────────────────────────────────────────────
        contratos = Contrato.objects.select_related("unidade_requisitante").all()
        total = contratos.count()
        vinculados = 0
        ignorados  = 0
        sem_match: list[tuple[str, str]] = []

        self.stdout.write(f"Processando {total} contratos...\n")

        for c in contratos:
            if c.unidade_requisitante and not force:
                ignorados += 1
                continue

            num_norm = _norm_numero(c.numero_contrato)
            setor, fonte = None, "—"

            # P1: número → setor direto (2020-2024)
            if num_setor.get(num_norm):
                setor, fonte = num_setor[num_norm], "direto"

            # P2: numero_sei do banco → setor
            if not setor and c.numero_sei:
                setor = sei_setor.get(_norm_sei(c.numero_sei))
                if setor:
                    fonte = "sei_banco"

            # P3: número → SEI (planilha) → setor
            if not setor:
                sei_p = num_sei.get(num_norm)
                if sei_p:
                    setor = sei_setor.get(sei_p)
                    if setor:
                        fonte = f"sei={sei_p}"

            # P4: número de NE como numero_contrato (2020-2022)
            if not setor:
                ne_key = c.numero_contrato.strip().upper()
                # Tratar múltiplas NEs separadas por ";"
                for ne in re.split(r"[;,]+", ne_key):
                    ne = ne.strip()
                    if ne in ne_setor:
                        setor, fonte = ne_setor[ne], "ne_direto"
                        break

            if not setor:
                sem_match.append((c.numero_contrato, c.numero_sei or ""))
                continue

            unidade = unidades.get(setor.upper())
            if not unidade:
                self.stdout.write(self.style.WARNING(
                    f"  Setor '{setor}' sem UnidadeRequisitante (contrato {c.numero_contrato})"
                ))
                continue

            vinculados += 1
            if dry_run:
                self.stdout.write(
                    f"  [DRY] {c.numero_contrato:<28} → {setor:<10} [{fonte}]"
                )
            else:
                c.unidade_requisitante = unidade
                c.save(update_fields=["unidade_requisitante"])
                self.stdout.write(f"  ✓ {c.numero_contrato:<28} → {setor}  [{fonte}]")

        # ── 6. Resumo ─────────────────────────────────────────────────────────
        self.stdout.write("\n" + "─" * 65)
        self.stdout.write(self.style.SUCCESS(
            f"{'[DRY-RUN] ' if dry_run else ''}Vinculados: {vinculados} / {total}"
        ))
        if ignorados:
            self.stdout.write(
                f"  Ignorados (vínculo existente — use --force para sobrescrever): {ignorados}"
            )
        if sem_match:
            self.stdout.write(self.style.WARNING(
                f"\nSem correspondência ({len(sem_match)}):"
            ))
            for num, sei in sem_match:
                self.stdout.write(f"  {num:<28}  sei_banco={sei or '—'}")

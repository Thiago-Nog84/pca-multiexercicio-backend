"""
Importa o Catalogo Interno 2027 (planilha PDM/CATSER) para ItemCatalogo.

Fonte: catalogo_interno_2027_PDM_CATSER_preenchido.xlsx (aba "catalogo_interno_2027")
Colunas: ORDEM | NOME | DESCRICAO | TIPO | PDM/CATSER | GRUPO | ATIVO | EXERCICIO | VALOR DO ITEM

Regras:
- Chave de idempotencia: (descricao_padrao=NOME, codigo_catmat_catser=PDM/CATSER).
  Re-execucoes atualizam em vez de duplicar. Duplicatas exatas na planilha sao puladas.
- codigo_catalogo gerado sequencialmente: CONT-FORN-NNN (Material) / CONT-SERV-NNN (Servico),
  continuando a numeracao ja existente no banco.
- GRUPO normalizado (variacoes de caixa como "Manutencao predial - Civil" vs
  "Manutencao Predial - Civil" sao unificadas pela primeira ocorrencia).
- ATIVO vazio e tratado como True (linhas nao preenchidas na planilha).
- Categoria inferida de TIPO + GRUPO (heuristica conservadora — revisar depois):
    Servico + grupo "Eventos, Cursos e Capacitacao"  -> treinamento
    grupo contendo "Software"                        -> software
    Material + grupo tipicamente permanente          -> material_permanente
    Material (demais)                                -> material
    Servico (demais)                                 -> servico

Uso:
    python manage.py importar_catalogo_interno --arquivo caminho/planilha.xlsx [--dry-run]
"""

import re
import unicodedata
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.pca.models import ItemCatalogo

GRUPOS_PERMANENTES = [
    "mobiliario",
    "eletroeletronico",
    "eletrodomestico",
    "informatica - equipamentos",
    "informatica - infraestrutura",
    "veiculos",
    "seguranca eletronica",
]


def _sem_acento(texto):
    return unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")


class Command(BaseCommand):
    help = "Importa a planilha do Catalogo Interno 2027 para ItemCatalogo."

    def add_arguments(self, parser):
        parser.add_argument("--arquivo", required=True, help="Caminho do .xlsx")
        parser.add_argument("--aba", default="catalogo_interno_2027")
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Simula sem gravar no banco.",
        )

    # ------------------------------------------------------------------
    def _categoria(self, tipo, grupo):
        g = _sem_acento(grupo or "").lower()
        if "software" in g:
            return "software"
        if tipo == "servico":
            if "eventos, cursos e capacitacao" in g:
                return "treinamento"
            return "servico"
        for perm in GRUPOS_PERMANENTES:
            if perm in g:
                return "material_permanente"
        return "material"

    def _proximo_codigo(self, prefixo, usados):
        n = usados.get(prefixo, 0) + 1
        usados[prefixo] = n
        return f"{prefixo}-{n:03d}"

    # ------------------------------------------------------------------
    def handle(self, *args, **opts):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl nao instalado: pip install openpyxl")

        try:
            wb = openpyxl.load_workbook(opts["arquivo"], read_only=True, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"Arquivo nao encontrado: {opts['arquivo']}")
        if opts["aba"] not in wb.sheetnames:
            raise CommandError(f"Aba '{opts['aba']}' nao existe. Abas: {wb.sheetnames}")

        ws = wb[opts["aba"]]
        linhas = list(ws.iter_rows(values_only=True))
        cabecalho = [str(c or "").strip().upper() for c in linhas[0]]
        esperado = ["ORDEM", "NOME", "DESCRICAO", "TIPO", "PDM/CATSER", "GRUPO"]
        if cabecalho[: len(esperado)] != esperado:
            raise CommandError(f"Cabecalho inesperado: {cabecalho}")

        # Continuar numeracao de codigos ja existente por prefixo
        usados = {}
        for cod in ItemCatalogo.objects.values_list("codigo_catalogo", flat=True):
            m = re.match(r"^(CONT-[A-Z]+)-(\d+)$", cod)
            if m:
                pref, num = m.group(1), int(m.group(2))
                usados[pref] = max(usados.get(pref, 0), num)

        grupos_canonicos = {}  # casefold -> forma canonica (primeira ocorrencia)
        vistos = set()         # (nome, codigo) ja processados nesta execucao
        criados = atualizados = pulados_dup = avisos = 0
        por_categoria = {}

        with transaction.atomic():
            for i, row in enumerate(linhas[1:], start=2):
                nome = str(row[1] or "").strip()
                if not nome:
                    continue
                descricao = str(row[2] or "").strip()
                tipo_raw = str(row[3] or "").strip().lower()
                codigo_pdm = str(row[4]).strip() if row[4] is not None else ""
                grupo_raw = str(row[5] or "").strip()
                ativo = row[6] is None or str(row[6]).strip().lower() == "true"
                valor = None
                if row[8] not in (None, ""):
                    try:
                        valor = Decimal(str(row[8]))
                    except InvalidOperation:
                        avisos += 1
                        self.stderr.write(f"  linha {i}: valor invalido {row[8]!r} — ignorado")

                if tipo_raw.startswith("material"):
                    tipo = "material"
                elif tipo_raw.startswith("servi"):
                    tipo = "servico"
                else:
                    avisos += 1
                    self.stderr.write(f"  linha {i}: TIPO vazio/desconhecido ({row[3]!r}) — assumindo material")
                    tipo = "material"

                if not codigo_pdm:
                    avisos += 1
                    self.stderr.write(f"  linha {i}: sem codigo PDM/CATSER ({nome[:50]})")

                # Grupo canonico (unifica variacoes de caixa)
                chave_g = _sem_acento(grupo_raw).casefold()
                if chave_g and chave_g not in grupos_canonicos:
                    grupos_canonicos[chave_g] = grupo_raw
                grupo = grupos_canonicos.get(chave_g, "")

                chave = (nome, codigo_pdm)
                if chave in vistos:
                    pulados_dup += 1
                    continue
                vistos.add(chave)

                categoria = self._categoria(tipo, grupo)
                por_categoria[categoria] = por_categoria.get(categoria, 0) + 1

                defaults = {
                    "descricao_detalhada": descricao,
                    "categoria": categoria,
                    "grupo": grupo,
                    "valor_referencia": valor,
                    "ativo": ativo,
                }
                obj = ItemCatalogo.objects.filter(
                    descricao_padrao=nome, codigo_catmat_catser=codigo_pdm
                ).first()
                if obj:
                    for campo, v in defaults.items():
                        setattr(obj, campo, v)
                    obj.save(update_fields=list(defaults))
                    atualizados += 1
                else:
                    prefixo = "CONT-FORN" if tipo == "material" else "CONT-SERV"
                    ItemCatalogo.objects.create(
                        codigo_catalogo=self._proximo_codigo(prefixo, usados),
                        descricao_padrao=nome,
                        codigo_catmat_catser=codigo_pdm,
                        **defaults,
                    )
                    criados += 1

            if opts["dry_run"]:
                transaction.set_rollback(True)

        modo = "[DRY-RUN — nada gravado] " if opts["dry_run"] else ""
        self.stdout.write(self.style.SUCCESS(
            f"{modo}Criados: {criados} | Atualizados: {atualizados} | "
            f"Duplicatas puladas: {pulados_dup} | Avisos: {avisos}"
        ))
        self.stdout.write("Distribuicao por categoria (heuristica — revisar):")
        for cat, qtd in sorted(por_categoria.items(), key=lambda x: -x[1]):
            self.stdout.write(f"  {cat}: {qtd}")
